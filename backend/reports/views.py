import csv
import io
from django.http import HttpResponse
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import permissions

from accounts.permissions import IsAnyStaff, IsUniversityCoordinator, scoped_dept_id
from reports.services.reporting import (
    course_offering_report, programme_report, department_report,
    semester_report, danger_report, graduation_report,
    university_comparison_report,
)
from reports.services.batch_upload import process_batch_upload
from reports.services.exports import course_pdf, department_pdf, to_excel
from results.models import CourseResult, ResultEditLog, SemesterResult


def _get_dept_id(request):
    user = request.user
    if user.has_dept_scope:
        return user.department_id
    return request.query_params.get('department_id') or request.data.get('department_id')


# ── Course Offering ──────────────────────────────────────────────────────────

class CourseOfferingReportView(APIView):
    permission_classes = [IsAnyStaff]
    def get(self, request):
        course_id = request.query_params.get('course_id')
        if not course_id:
            return Response({'error': 'course_id required'}, status=400)
        data = course_offering_report(
            course_id=int(course_id),
            ay_id=request.query_params.get('academic_year_id'),
            semester=request.query_params.get('semester'),
            year_of_study=request.query_params.get('year_of_study'),
            dept_id=_get_dept_id(request),
            order_by=request.query_params.get('order_by', 'index'),
            descending=request.query_params.get('desc', '').lower() == 'true',
        )
        return Response(data)


class CourseOfferingPDFView(APIView):
    permission_classes = [IsAnyStaff]
    def get(self, request):
        course_id = request.query_params.get('course_id')
        if not course_id:
            return Response({'error': 'course_id required'}, status=400)
        data = course_offering_report(
            course_id=int(course_id),
            ay_id=request.query_params.get('academic_year_id'),
            semester=request.query_params.get('semester'),
            year_of_study=request.query_params.get('year_of_study'),
            dept_id=_get_dept_id(request),
        )
        buf = course_pdf(data, generated_by=request.user.get_full_name())
        code = data.get('course', {}).get('code', 'course')
        resp = HttpResponse(buf, content_type='application/pdf')
        resp['Content-Disposition'] = f'attachment; filename="course_{code}.pdf"'
        return resp


class CourseOfferingExcelView(APIView):
    permission_classes = [IsAnyStaff]
    def get(self, request):
        course_id = request.query_params.get('course_id')
        if not course_id:
            return Response({'error': 'course_id required'}, status=400)
        data = course_offering_report(
            course_id=int(course_id),
            ay_id=request.query_params.get('academic_year_id'),
            semester=request.query_params.get('semester'),
            dept_id=_get_dept_id(request),
        )
        headers = ['#','Index Number','Full Name','Department','Programme',
                   'Year','CA','Exam','Total','Grade','Pass/Fail']
        rows = [{'#':i,'index_number':r['index_number'],'full_name':r['full_name'],
                 'department':r['department'],'programme':r['programme'],'year':r['year_of_study'],
                 'ca':r['ca'],'exam':r['exam'],'total':r['total'],'grade':r['grade'],
                 'pass/fail':'Pass' if r['pass'] else 'Fail'}
                for i,r in enumerate(data['rows'],1)]
        code = data.get('course',{}).get('code','course')
        buf = to_excel(headers, rows, sheet_title=f'{code} Report',
                       summary_data={k:v for k,v in data['analytics'].items() if not isinstance(v,dict)})
        resp = HttpResponse(buf, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        resp['Content-Disposition'] = f'attachment; filename="course_{code}.xlsx"'
        return resp


class CourseOfferingCSVView(APIView):
    permission_classes = [IsAnyStaff]
    def get(self, request):
        course_id = request.query_params.get('course_id')
        if not course_id:
            return Response({'error': 'course_id required'}, status=400)
        data = course_offering_report(course_id=int(course_id), dept_id=_get_dept_id(request),
                                       ay_id=request.query_params.get('academic_year_id'),
                                       semester=request.query_params.get('semester'))
        resp = HttpResponse(content_type='text/csv')
        resp['Content-Disposition'] = 'attachment; filename="course_report.csv"'
        fields = ['index_number','full_name','department','programme','year_of_study','ca','exam','total','grade','pass']
        writer = csv.DictWriter(resp, fieldnames=fields, extrasaction='ignore')
        writer.writeheader()
        writer.writerows(data['rows'])
        return resp


# ── Programme Report ─────────────────────────────────────────────────────────

class ProgrammeReportView(APIView):
    permission_classes = [IsAnyStaff]
    def get(self, request):
        prog_id = request.query_params.get('programme_id')
        if not prog_id:
            return Response({'error': 'programme_id required'}, status=400)
        data = programme_report(
            programme_id=int(prog_id),
            dept_id=_get_dept_id(request),
            year=request.query_params.get('year'),
            ay_id=request.query_params.get('academic_year_id'),
            min_gpa=request.query_params.get('min_gpa'),
            max_gpa=request.query_params.get('max_gpa'),
        )
        return Response(data)


class ProgrammeExcelView(APIView):
    permission_classes = [IsAnyStaff]
    def get(self, request):
        prog_id = request.query_params.get('programme_id')
        if not prog_id:
            return Response({'error': 'programme_id required'}, status=400)
        data = programme_report(
            programme_id=int(prog_id), dept_id=_get_dept_id(request),
            year=request.query_params.get('year'),
            min_gpa=request.query_params.get('min_gpa'),
            max_gpa=request.query_params.get('max_gpa'),
        )
        headers = ['#','Index Number','Full Name','Department','Programme',
                   'Year','Status','CuGPA','Standing','Trails','In Danger']
        rows = [{'#':i,'index_number':r['index_number'],'full_name':r['full_name'],
                 'department':r['department'],'programme':r['programme'],'year':r['current_year'],
                 'status':r['status'],'cugpa':r['cumulative_gpa'],'standing':r['academic_standing'],
                 'trails':r['trail_count'],'in_danger':'Yes' if r['in_danger'] else 'No'}
                for i,r in enumerate(data['rows'],1)]
        buf = to_excel(headers, rows, sheet_title='Programme Report',
                       summary_data={k:v for k,v in data['summary'].items() if not isinstance(v,dict)})
        resp = HttpResponse(buf, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        resp['Content-Disposition'] = 'attachment; filename="programme_report.xlsx"'
        return resp


# ── Department Report ────────────────────────────────────────────────────────

class DepartmentReportView(APIView):
    permission_classes = [IsAnyStaff]
    def get(self, request):
        dept_id = _get_dept_id(request)
        if not dept_id:
            return Response({'error': 'department_id required'}, status=400)
        data = department_report(
            dept_id=int(dept_id),
            programme_id=request.query_params.get('programme_id'),
            year=request.query_params.get('year'),
            standing_filter=request.query_params.get('standing'),
            min_gpa=request.query_params.get('min_gpa'),
            max_gpa=request.query_params.get('max_gpa'),
            order_by=request.query_params.get('order_by', 'name'),
            descending=request.query_params.get('desc', '').lower() == 'true',
        )
        return Response(data)

class DepartmentReportPDFView(APIView):
    permission_classes = [IsAnyStaff]
    def get(self, request):
        dept_id = _get_dept_id(request)
        if not dept_id:
            return Response({'error': 'department_id required'}, status=400)
        data = department_report(dept_id=int(dept_id),
                                  year=request.query_params.get('year'),
                                  standing_filter=request.query_params.get('standing'))
        buf = department_pdf(data, generated_by=request.user.get_full_name())
        code = data.get('department',{}).get('code','dept')
        resp = HttpResponse(buf, content_type='application/pdf')
        resp['Content-Disposition'] = f'attachment; filename="dept_{code}.pdf"'
        return resp


class DepartmentReportExcelView(APIView):
    permission_classes = [IsAnyStaff]
    def get(self, request):
        dept_id = _get_dept_id(request)
        if not dept_id:
            return Response({'error': 'department_id required'}, status=400)
        data = department_report(dept_id=int(dept_id),
                                  year=request.query_params.get('year'),
                                  standing_filter=request.query_params.get('standing'),
                                  min_gpa=request.query_params.get('min_gpa'),
                                  max_gpa=request.query_params.get('max_gpa'))
        headers = ['#','Index Number','Full Name','Programme','Year',
                   'Status','CuGPA','Standing','Trails','In Danger']
        rows = [{'#':i,'index_number':r['index_number'],'full_name':r['full_name'],
                 'programme':r['programme'],'year':r['current_year'],'status':r['status'],
                 'cugpa':r['cumulative_gpa'],'standing':r['academic_standing'],
                 'trails':r['trail_count'],'in_danger':'Yes' if r['in_danger'] else 'No'}
                for i,r in enumerate(data['rows'],1)]
        buf = to_excel(headers, rows, sheet_title='Department Report',
                       summary_data={k:v for k,v in data['summary'].items() if not isinstance(v,dict)})
        resp = HttpResponse(buf, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        resp['Content-Disposition'] = 'attachment; filename="department_report.xlsx"'
        return resp


# ── Semester Report ──────────────────────────────────────────────────────────

class SemesterReportView(APIView):
    permission_classes = [IsAnyStaff]
    def get(self, request):
        ay_id = request.query_params.get('academic_year_id')
        semester = request.query_params.get('semester')
        if not ay_id or not semester:
            return Response({'error': 'academic_year_id and semester required'}, status=400)
        data = semester_report(
            ay_id=int(ay_id), semester=int(semester),
            dept_id=_get_dept_id(request),
            programme_id=request.query_params.get('programme_id'),
            year_of_study=request.query_params.get('year_of_study'),
        )
        return Response(data)


class SemesterReportExcelView(APIView):
    permission_classes = [IsAnyStaff]
    def get(self, request):
        ay_id = request.query_params.get('academic_year_id')
        semester = request.query_params.get('semester')
        if not ay_id or not semester:
            return Response({'error': 'academic_year_id and semester required'}, status=400)
        data = semester_report(ay_id=int(ay_id), semester=int(semester),
                                dept_id=_get_dept_id(request))
        headers = ['#','Index Number','Full Name','Department','Programme',
                   'Year','Semester GPA','Credits','CuGPA','In Danger']
        rows = [{'#':i,'index_number':r['index_number'],'full_name':r['full_name'],
                 'department':r['department'],'programme':r['programme'],'year':r['year_of_study'],
                 'semester_gpa':r['semester_gpa'],'credits':r['total_credits'],
                 'cugpa':r['cumulative_gpa'],'in_danger':'Yes' if r['in_danger'] else 'No'}
                for i,r in enumerate(data['rows'],1)]
        buf = to_excel(headers, rows, sheet_title='Semester Report', summary_data=data['summary'])
        resp = HttpResponse(buf, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        resp['Content-Disposition'] = 'attachment; filename="semester_report.xlsx"'
        return resp


# ── Danger Report ────────────────────────────────────────────────────────────

class DangerReportView(APIView):
    permission_classes = [IsAnyStaff]
    def get(self, request):
        data = danger_report(
            dept_id=_get_dept_id(request),
            programme_id=request.query_params.get('programme_id'),
            year=request.query_params.get('year'),
            gpa_threshold=request.query_params.get('gpa_threshold', 1.5),
            min_trails=request.query_params.get('min_trails', 3),
        )
        return Response(data)


class DangerReportExcelView(APIView):
    permission_classes = [IsAnyStaff]
    def get(self, request):
        data = danger_report(
            dept_id=_get_dept_id(request),
            gpa_threshold=request.query_params.get('gpa_threshold', 1.5),
            min_trails=request.query_params.get('min_trails', 3),
        )
        headers = ['#','Index Number','Full Name','Department','Programme',
                   'Year','CuGPA','Standing','Trails','GPA Risk','Trail Risk']
        rows = [{'#':i,'index_number':r['index_number'],'full_name':r['full_name'],
                 'department':r['department'],'programme':r['programme'],'year':r['current_year'],
                 'cugpa':r['cumulative_gpa'],'standing':r['academic_standing'],
                 'trails':r['trail_count'],
                 'gpa_risk':'Yes' if r['gpa_below_threshold'] else 'No',
                 'trail_risk':'Yes' if r['trails_exceeded'] else 'No'}
                for i,r in enumerate(data['rows'],1)]
        buf = to_excel(headers, rows, sheet_title='Academic Danger', summary_data=data['summary'])
        resp = HttpResponse(buf, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        resp['Content-Disposition'] = 'attachment; filename="danger_report.xlsx"'
        return resp


# ── Graduation Report ────────────────────────────────────────────────────────

class GraduationReportView(APIView):
    permission_classes = [IsAnyStaff]
    def get(self, request):
        data = graduation_report(
            dept_id=_get_dept_id(request),
            programme_id=request.query_params.get('programme_id'),
            min_credits=request.query_params.get('min_credits', 120),
        )
        return Response(data)


class GraduationReportExcelView(APIView):
    permission_classes = [IsAnyStaff]
    def get(self, request):
        data = graduation_report(
            dept_id=_get_dept_id(request),
            programme_id=request.query_params.get('programme_id'),
            min_credits=request.query_params.get('min_credits', 120),
        )
        headers = ['#','Index Number','Full Name','Programme',
                   'CuGPA','Standing','Credits Earned','Trails','Eligibility']
        rows = [{'#':i,'index_number':r['index_number'],'full_name':r['full_name'],
                 'programme':r['programme'],'cugpa':r['cumulative_gpa'],
                 'standing':r['academic_standing'],'credits_earned':r['credits_earned'],
                 'trails':r['trail_count'],'eligibility':r['eligibility']}
                for i,r in enumerate(data['rows'],1)]
        buf = to_excel(headers, rows, sheet_title='Graduation Eligibility', summary_data=data['summary'])
        resp = HttpResponse(buf, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        resp['Content-Disposition'] = 'attachment; filename="graduation_report.xlsx"'
        return resp


# ── University Comparison ────────────────────────────────────────────────────

class UniversityComparisonView(APIView):
    permission_classes = [IsUniversityCoordinator]
    def get(self, request):
        return Response(university_comparison_report())


# ── Batch Upload ─────────────────────────────────────────────────────────────

class BatchUploadView(APIView):
    permission_classes = [IsAnyStaff]
    def post(self, request):
        file = request.FILES.get('file')
        if not file:
            return Response({'error': 'No file uploaded.'}, status=400)
        result = process_batch_upload(
            file_bytes=file.read(),
            filename=file.name,
            entered_by=request.user,
            dept_id=_get_dept_id(request),
        )
        return Response(result)


# ── Result Edit + Audit Log ──────────────────────────────────────────────────

class ResultEditView(APIView):
    permission_classes = [IsAnyStaff]
    def patch(self, request, pk):
        try:
            cr = CourseResult.objects.select_related('student__department').get(pk=pk)
        except CourseResult.DoesNotExist:
            return Response({'error': 'Result not found'}, status=404)

        user = request.user
        if user.has_dept_scope and cr.student.department_id != user.department_id:
            return Response({'error': 'Forbidden: outside your department'}, status=403)
        if cr.is_locked and not user.is_university_coordinator:
            return Response({'error': 'Result is locked. Only University Coordinator can override.'}, status=403)

        old_ca, old_exam, old_grade = cr.continuous_assessment, cr.exam_score, cr.grade
        new_ca   = request.data.get('ca')
        new_exam = request.data.get('exam')
        if new_ca   is not None: cr.continuous_assessment = float(new_ca)
        if new_exam is not None: cr.exam_score = float(new_exam)
        cr.save()

        ResultEditLog.objects.create(
            course_result=cr, edited_by=user,
            old_ca=old_ca, new_ca=cr.continuous_assessment,
            old_exam=old_exam, new_exam=cr.exam_score,
            old_grade=old_grade, new_grade=cr.grade,
            reason=request.data.get('reason', ''),
            is_override=cr.is_locked,
        )

        try:
            sr = SemesterResult.objects.get(
                student=cr.student, academic_year=cr.academic_year, semester=cr.semester)
            sr.recompute()
        except SemesterResult.DoesNotExist:
            pass

        return Response({'id': cr.id, 'grade': cr.grade, 'total_score': cr.total_score,
                         'grade_point': cr.grade_point, 'message': 'Updated and logged.'})


class ResultLockView(APIView):
    permission_classes = [IsAnyStaff]
    def patch(self, request, pk):
        try:
            cr = CourseResult.objects.get(pk=pk)
        except CourseResult.DoesNotExist:
            return Response({'error': 'Not found'}, status=404)
        lock = request.data.get('is_locked', True)
        if not lock and not request.user.is_university_coordinator:
            return Response({'error': 'Only University Coordinator can unlock.'}, status=403)
        cr.is_locked = lock
        cr.save(update_fields=['is_locked'])
        return Response({'is_locked': cr.is_locked})


class ResultEditLogView(APIView):
    permission_classes = [IsAnyStaff]
    def get(self, request, pk):
        logs = ResultEditLog.objects.filter(
            course_result_id=pk).select_related('edited_by').order_by('-timestamp')
        return Response([{
            'id': l.id, 'edited_by': l.edited_by.get_full_name() if l.edited_by else '—',
            'old_ca': l.old_ca, 'new_ca': l.new_ca,
            'old_exam': l.old_exam, 'new_exam': l.new_exam,
            'old_grade': l.old_grade, 'new_grade': l.new_grade,
            'reason': l.reason, 'is_override': l.is_override,
            'timestamp': l.timestamp,
        } for l in logs])


class HODDashboardView(APIView):
    permission_classes = [IsAnyStaff]

    def get(self, request):
        from accounts.permissions import scoped_dept_id
        from students.models import Student
        from results.models import ElectivePool
        from registration.models import CourseRegistration
        from grievances.models import Grievance
        from django.utils import timezone

        dept_id = scoped_dept_id(request.user)
        students = Student.objects.filter(department_id=dept_id) if dept_id else Student.objects.all()
        regs = CourseRegistration.objects.filter(student__department_id=dept_id) if dept_id \
               else CourseRegistration.objects.all()
        pools = ElectivePool.objects.filter(department_id=dept_id, is_active=True) if dept_id \
                else ElectivePool.objects.filter(is_active=True)
        grievances = Grievance.objects.filter(student__department_id=dept_id) if dept_id \
                     else Grievance.objects.all()

        danger_count = sum(1 for s in students if s.in_academic_danger)

        return Response({
            'students': {
                'total':    students.count(),
                'active':   students.filter(status='active').count(),
                'repeating':students.filter(status='repeating').count(),
                'deferred': students.filter(status='deferred').count(),
                'danger':   danger_count,
            },
            'registrations': {
                'total':    regs.count(),
                'draft':    regs.filter(status='draft').count(),
                'submitted':regs.filter(status='submitted').count(),
                'approved': regs.filter(status='approved').count(),
                'rejected': regs.filter(status='rejected').count(),
            },
            'elective_pools': [{
                'id': p.id, 'name': p.name,
                'year': p.year_of_study, 'semester': p.semester,
                'course_count': p.pool_courses.count(),
                'max_electives': p.max_electives,
            } for p in pools],
            'grievances': {
                'total':   grievances.count(),
                'pending': grievances.filter(status='pending').count(),
                'high':    grievances.filter(priority='high',
                               status__in=['pending','in_review','escalated']).count(),
                'overdue': grievances.filter(
                               response_deadline__lt=timezone.now()
                           ).exclude(status__in=['resolved','rejected']).count(),
            },
        })
class CourseRegistrationStatusView(APIView):
    """
    For each course in the department, show which students are registered
    and which are not. Accessible by HoD and dept coordinator.
    """
    permission_classes = [IsAnyStaff]

    def get(self, request):
        from accounts.permissions import scoped_dept_id
        from results.models import Course
        from students.models import Student
        from registration.models import RegisteredCourse

        dept_id = scoped_dept_id(request.user)
        if not dept_id:
            return Response({'error': 'Department required.'}, status=400)

        # Optional filters
        year     = request.query_params.get('year')
        semester = request.query_params.get('semester')
        course_id= request.query_params.get('course_id')

        # All active courses in the department
        courses_qs = Course.objects.filter(
            department_id=dept_id, is_active=True
        ).order_by('year', 'semester', 'code')

        if year:     courses_qs = courses_qs.filter(year=int(year))
        if semester: courses_qs = courses_qs.filter(semester=int(semester))
        if course_id:courses_qs = courses_qs.filter(id=int(course_id))

        # All students in the department
        all_students = list(
            Student.objects.filter(department_id=dept_id)
            .select_related('user', 'programme')
            .order_by('user__last_name', 'user__first_name')
        )

        result = []
        for course in courses_qs:
            # Students registered for this course (any registration status)
            registered_entries = RegisteredCourse.objects.filter(
                course=course,
                registration__student__department_id=dept_id,
            ).select_related(
                'registration__student__user',
                'registration__student__programme',
            )

            registered_student_ids = set(
                rc.registration.student_id for rc in registered_entries
            )

            registered = []
            for rc in registered_entries:
                s = rc.registration.student
                registered.append({
                    'student_id':    s.id,
                    'index_number':  s.index_number,
                    'full_name':     s.user.get_full_name(),
                    'programme':     s.programme.name,
                    'year':          s.current_year,
                    'reg_status':    rc.registration.status,
                    'is_core':       rc.is_core,
                })

            not_registered = []
            for s in all_students:
                # Only flag students whose year matches the course year
                if s.current_year != course.year:
                    continue
                if s.id not in registered_student_ids:
                    not_registered.append({
                        'student_id':   s.id,
                        'index_number': s.index_number,
                        'full_name':    s.user.get_full_name(),
                        'programme':    s.programme.name,
                        'year':         s.current_year,
                    })

            result.append({
                'course_id':        course.id,
                'code':             course.code,
                'title':            course.title,
                'credit_hours':     course.credit_hours,
                'year':             course.year,
                'semester':         course.semester,
                'is_core':          course.is_core,
                'registered_count': len(registered),
                'not_registered_count': len(not_registered),
                'registered':       registered,
                'not_registered':   not_registered,
            })

        return Response({'courses': result, 'department_id': dept_id})


class CourseRegistrationStatusExcelView(APIView):
    """Download the registration status as Excel."""
    permission_classes = [IsAnyStaff]

    def get(self, request):
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from accounts.permissions import scoped_dept_id
        from results.models import Course
        from students.models import Student
        from registration.models import RegisteredCourse
        import io

        dept_id = scoped_dept_id(request.user)
        if not dept_id:
            return Response({'error': 'Department required.'}, status=400)

        year     = request.query_params.get('year')
        semester = request.query_params.get('semester')
        course_id= request.query_params.get('course_id')

        courses_qs = Course.objects.filter(
            department_id=dept_id, is_active=True
        ).order_by('year', 'semester', 'code')
        if year:      courses_qs = courses_qs.filter(year=int(year))
        if semester:  courses_qs = courses_qs.filter(semester=int(semester))
        if course_id: courses_qs = courses_qs.filter(id=int(course_id))

        all_students = list(
            Student.objects.filter(department_id=dept_id)
            .select_related('user', 'programme')
            .order_by('user__last_name', 'user__first_name')
        )

        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # remove default sheet

        NAVY   = 'FF0F2744'
        GREEN  = 'FF16a34a'
        RED    = 'FFdc2626'
        GRAY   = 'FFF1F5F9'

        for course in courses_qs:
            sheet_name = f"{course.code}"[:31]
            ws = wb.create_sheet(title=sheet_name)

            # Header
            ws.merge_cells('A1:F1')
            ws['A1'] = f"{course.code} — {course.title}"
            ws['A1'].font = Font(bold=True, size=13, color='FFFFFFFF')
            ws['A1'].fill = PatternFill('solid', fgColor=NAVY)
            ws['A1'].alignment = Alignment(horizontal='center')

            ws['A2'] = f"Year {course.year}  |  Semester {course.semester}  |  {course.credit_hours} credits  |  {'Core' if course.is_core else 'Elective'}"
            ws['A2'].font = Font(italic=True, color='FF555555')
            ws.merge_cells('A2:F2')

            # Registered section
            ws['A4'] = 'REGISTERED STUDENTS'
            ws['A4'].font = Font(bold=True, color='FF16a34a')

            reg_entries = RegisteredCourse.objects.filter(
                course=course, registration__student__department_id=dept_id
            ).select_related('registration__student__user','registration__student__programme')

            registered_ids = set()
            headers = ['#', 'Index Number', 'Full Name', 'Programme', 'Year', 'Reg. Status']
            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=5, column=col, value=h)
                cell.font = Font(bold=True, color='FFFFFFFF')
                cell.fill = PatternFill('solid', fgColor='FF16a34a')

            for i, rc in enumerate(reg_entries, 1):
                s = rc.registration.student
                registered_ids.add(s.id)
                ws.cell(row=5+i, column=1, value=i)
                ws.cell(row=5+i, column=2, value=s.index_number)
                ws.cell(row=5+i, column=3, value=s.user.get_full_name())
                ws.cell(row=5+i, column=4, value=s.programme.name)
                ws.cell(row=5+i, column=5, value=s.current_year)
                ws.cell(row=5+i, column=6, value=rc.registration.status)
                if i % 2 == 0:
                    for col in range(1, 7):
                        ws.cell(row=5+i, column=col).fill = PatternFill('solid', fgColor=GRAY)

            not_reg_row = 5 + reg_entries.count() + 3
            ws.cell(row=not_reg_row-1, column=1, value='NOT REGISTERED').font = Font(bold=True, color='FFdc2626')

            for col, h in enumerate(headers[:5], 1):
                cell = ws.cell(row=not_reg_row, column=col, value=h)
                cell.font = Font(bold=True, color='FFFFFFFF')
                cell.fill = PatternFill('solid', fgColor='FFdc2626')

            i = 1
            for s in all_students:
                if s.current_year != course.year: continue
                if s.id in registered_ids: continue
                ws.cell(row=not_reg_row+i, column=1, value=i)
                ws.cell(row=not_reg_row+i, column=2, value=s.index_number)
                ws.cell(row=not_reg_row+i, column=3, value=s.user.get_full_name())
                ws.cell(row=not_reg_row+i, column=4, value=s.programme.name)
                ws.cell(row=not_reg_row+i, column=5, value=s.current_year)
                i += 1

            for col in range(1, 7):
                ws.column_dimensions[chr(64+col)].width = [6,18,24,28,8,14][col-1]

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        resp = HttpResponse(
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        resp['Content-Disposition'] = 'attachment; filename="registration_status.xlsx"'
        return resp